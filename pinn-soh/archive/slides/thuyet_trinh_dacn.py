"""
Bộ slide thuyết trình đồ án chuyên ngành (bản cuối) — dựng trên khuôn BK của bộ slide báo cáo kết quả 30/09.

    python slide/thuyet_trinh_dacn.py      # -> slide/Thuyet_trinh_DACN_Nhom15.pptx và .pdf

Chạy sau `python run.py analyze` (cần results/fig_*.pdf, results/e30_metrics.csv) và sau khi dựng báo cáo
(do-an/main.tex): mọi con số viết trong slide được đối chiếu với lời văn của báo cáo, bảng MAE/R² đọc thẳng từ
results/. Mỗi phần đã làm một slide có hình của báo cáo; chữ trên slide ngắn, lời nói nằm trong ghi chú.
"""
from __future__ import annotations
import os, re, subprocess, sys, tempfile
import pandas as pd
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import pptx_tools as T

SRC = os.path.join(HERE, 'Bao_cao_ket_qua_SOH_Nhom15.pptx')
DST = os.path.join(HERE, 'Thuyet_trinh_DACN_Nhom15.pptx')
R = os.path.join(ROOT, 'results')
TEX = open(os.path.join(ROOT, 'do-an', 'main.tex'), encoding='utf-8').read()
NAVY, INK, GREEN = '20567B', '163047', '0D6E56'
COVER, TEXT, TABLE, FRAME = 0, 1, 2, 7            # các slide mẫu trong bộ 30/09


def chk(*phrases):
    """Mỗi cụm số viết trên slide phải có trong báo cáo (đổi '0,39–0,90' thành '0{,}39--0{,}90')."""
    for p in phrases:
        t = p.replace('–', '--')
        t = re.sub(r'(\d),(\d)', r'\1{,}\2', t)
        assert t in TEX, f'không thấy trong báo cáo: {p!r} ({t!r})'
    return phrases[0] if len(phrases) == 1 else phrases


# ─────────────────────────────────────────────────────────────── hình
_PNG = tempfile.mkdtemp(prefix='slide_hinh_')


def png(name, dpi=250):
    """results/<name>.pdf -> PNG tạm (cùng hình của báo cáo)."""
    out = os.path.join(_PNG, name)
    if not os.path.exists(out + '.png'):
        subprocess.run(['pdftoppm', '-r', str(dpi), '-png', '-singlefile', os.path.join(R, name + '.pdf'), out],
                       check=True, stderr=subprocess.DEVNULL)
    return out + '.png'


# ─────────────────────────────────────────────────────────────── khuôn
prs = Presentation(SRC)
n_old = len(prs.slides)
_tpl_text = prs.slides[TEXT]
P_HEAD = [s for s in _tpl_text.shapes if s.has_text_frame and s.text_frame.text == 'Đầu vào và đầu ra'][0]
P_BODY = [s for s in _tpl_text.shapes if s.has_text_frame and s.text_frame.text.startswith('16 đặc trưng')][0]
P_FRAME = None
for _s in prs.slides[FRAME].shapes:              # khung viền nhạt quanh hình (hình chữ nhật không chữ, y ≈ 1,4 in)
    if _s.shape_type == 1 and Inches(1.2) < _s.top < Inches(1.6) and not _s.text_frame.text.strip():
        P_FRAME = _s
P_TABLE = [s for s in prs.slides[TABLE].shapes if s.has_table][0]
assert P_FRAME is not None


def new_slide(title):
    """Slide nội dung trống: giữ dải tiêu đề, số trang, logo; bỏ mọi khối khác của slide mẫu."""
    s = T.dup_slide(prs, TEXT)
    keep = {T.title_shape(s).shape_id, T.page_shape(s).shape_id}
    for sh in list(s.shapes):
        is_bar = sh.shape_type == 1 and sh.top < Inches(0.1) and sh.height > Inches(0.9) and not sh.text_frame.text.strip()
        is_logo = sh.shape_type == 13 and sh.top > Inches(6.8)
        if sh.shape_id not in keep and not is_bar and not is_logo:
            T.remove_shape(sh)
    T.set_text(T.title_shape(s), title)
    return s


def _clone(proto, s, x, y, w, h):
    sh = T.clone_shape(proto, s)
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)
    return sh


def heading(s, x, y, w, text, size=None):
    sh = T.set_text(_clone(P_HEAD, s, x, y, w, 0.36), text)
    if size:
        T.set_font_size(sh, size)
    return sh


def body(s, x, y, w, h, text, size=16, bullets=False):
    """text: chuỗi (xuống dòng = đoạn mới) hoặc danh sách đoạn; bullets=True: chấm tròn treo đầu dòng."""
    sh = T.set_text(_clone(P_BODY, s, x, y, w, h), text)
    T.set_font_size(sh, size)
    sh.text_frame.vertical_anchor = MSO_ANCHOR.TOP          # khối chữ bám đầu khung, ngay dưới đầu đề
    for p in sh.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        for tag in ('a:spcAft', 'a:buNone', 'a:buChar', 'a:buFont', 'a:buAutoNum'):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        sp = etree.Element(qn('a:spcAft'))
        etree.SubElement(sp, qn('a:spcPts')).set('val', str(int(size * 45)))
        new = [sp]
        if bullets:
            pPr.set('marL', str(int(Inches(0.27)))); pPr.set('indent', str(-int(Inches(0.27))))
            bf = etree.Element(qn('a:buFont')); bf.set('typeface', 'Arial')
            bc = etree.Element(qn('a:buChar')); bc.set('char', '•')
            new += [bf, bc]
        _insert_ordered(pPr, new)
    return sh


# thứ tự con của a:pPr theo lược đồ DrawingML (PowerPoint từ chối tệp sai thứ tự)
_PPR_ORDER = ['lnSpc', 'spcBef', 'spcAft', 'buClrTx', 'buClr', 'buSzTx', 'buSzPct', 'buSzPts', 'buFontTx', 'buFont',
              'buNone', 'buAutoNum', 'buChar', 'buBlip', 'tabLst', 'defRPr', 'extLst']


def _insert_ordered(pPr, new):
    kids = list(pPr) + new
    for k in list(pPr):
        pPr.remove(k)
    rank = lambda e: _PPR_ORDER.index(etree.QName(e).localname) if etree.QName(e).localname in _PPR_ORDER else 99
    for k in sorted(kids, key=rank):
        pPr.append(k)


def block(s, x, y, w, items, size=16, gap=0.12, line_h=None):
    """Chuỗi (đầu đề, lời) xếp dọc; trả y kế tiếp. Chiều cao lời ước theo số ký tự."""
    lh = line_h or size / 72 * 1.25
    for head, text in items:
        heading(s, x, y, w, head, size=size + 2)
        y += 0.42
        n = sum(max(1, -(-len(par) // int(w * 136 / size))) for par in text.split('\n'))
        h = n * lh + 0.08
        body(s, x, y, w, h, text, size=size)
        y += h + gap
    return y


def figure(s, name, x, y, w, h, frame=True, pad=0.12):
    if frame:
        _clone(P_FRAME, s, x, y, w, h)
    return T.add_picture_fit(s, png(name), Inches(x + pad), Inches(y + pad), Inches(w - 2 * pad), Inches(h - 2 * pad))


def table(s, x, y, header, rows, widths, font=13, row_h=0.42, best=(), bold_cols=()):
    """best: các ô (hàng, cột) của thân bảng in đậm màu xanh."""
    sh = _clone(P_TABLE, s, x, y, sum(widths), 1)
    T.fill_table(sh, rows, header=header, col_widths_in=widths, font_pt=font, row_h_in=row_h, bold_cols=bold_cols)
    for i, j in best:
        for r in sh.table.cell(i + 1, j).text_frame.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor.from_string(GREEN)
    return sh


def notes(s, text):
    T.set_notes(s, text)


# ═══════════════════════════════════════════════════════════ 1 · bìa
cov = prs.slides[COVER]
for sh in cov.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith('Huỳnh Quốc Thắng'):
        T.set_text(sh, 'Huỳnh Quốc Thắng · MSSV 2313178 · Nhóm 15')
    elif t.startswith('Báo cáo kết quả'):
        T.set_text(sh, ['Thuyết trình đồ án chuyên ngành · Khoa học Máy tính', 'TP. Hồ Chí Minh, 10/2026'])
notes(cov, 'Em xin trình bày đồ án chuyên ngành: dự đoán thông số sức khoẻ của pin bằng mô hình học sâu có thông tin '
           'hoá lý. Bài nói đi theo đúng các phần của đề cương giai đoạn 1, mỗi phần có hình kết quả.')

# ═══════════════════════════════════════════════════════════ 2 · nội dung
s = new_slide('Nội dung')
parts = ['Bài toán và mục tiêu theo đề cương', 'Dữ liệu: bốn bộ công khai và NASA', 'Mô hình PINN bán giám sát',
         'Hiện thực hệ thống và giao thức đánh giá', 'Kết quả theo lượng nhãn',
         'So với MLP, CNN, LSTM, Transformer; bảng MAE và R²', 'Mô hình kết hợp PINN-lai, đầu vào có lịch sử',
         'So với PINN4SOH', 'Dự báo SOH, dung lượng, RUL', 'Bộ NASA, SambaMixer, mức đo', 'Ngân sách đo dung lượng',
         'Đối chiếu đề cương, kết luận, giai đoạn 2']
for i, p in enumerate(parts):
    col, row = divmod(i, 6)
    x, y = 0.75 + col * 6.2, 1.55 + row * 0.86
    heading(s, x, y, 0.7, f'{i + 1:02d}', size=22)
    body(s, x + 0.8, y + 0.02, 5.2, 0.5, p, size=17)
notes(s, 'Mười hai phần, đi từ bài toán tới kết quả và phần đối chiếu với đề cương.')

# ═══════════════════════════════════════════════════════════ 3 · bài toán
s = new_slide('Bài toán và mục tiêu theo đề cương')
y = block(s, 0.68, 1.45, 11.9, [
    ('Bài toán', 'Ước lượng SOH = Q / Q danh định của từng cell ở từng chu kỳ từ 16 đặc trưng của đoạn sạc CC–CV, rồi dự báo '
                 'SOH, dung lượng và RUL. Nhãn đắt: muốn biết dung lượng phải xả đầy cell, còn đặc trưng sạc thì BMS '
                 'ghi được ở mọi chu kỳ.'),
], size=16)
heading(s, 0.68, y, 11.9, 'Mục tiêu giai đoạn 1', size=18)
body(s, 0.68, y + 0.42, 11.9, 2.4, [
    'Kiến thức nền: học sâu cho chuỗi thời gian, dữ liệu pin và các ràng buộc hoá lý',
    'Thống kê mô tả; phân tích phân phối, ngoại lai và mẫu gây nhiễu',
    'Chạy và so sánh ít nhất bốn mô hình nền (MLP, CNN, LSTM, Transformer) bằng MAE, RMSE, MAPE',
    'Tinh chỉnh mô hình, giải thích vì sao kết quả tốt lên hay xấu đi',
    'Đề xuất cải tiến có cơ sở lý thuyết cho giai đoạn 2'], size=16, bullets=True)
heading(s, 0.68, y + 2.35, 11.9, 'Hướng của đồ án', size=18)
body(s, 0.68, y + 2.77, 11.9, 0.7, 'PINN bán giám sát: ràng buộc vật lý không dùng nhãn được áp lên cả các cell chưa đo '
                                   'dung lượng, rồi kết hợp với mô hình cây.', size=16)
notes(s, 'Điểm mấu chốt là nhãn đắt: đo dung lượng phải xả đầy cell. Vì vậy đồ án tìm cách dùng dữ liệu sạc của các '
         'cell chưa đo, qua các ràng buộc vật lý không cần nhãn.')

# ═══════════════════════════════════════════════════════════ 4 · dữ liệu
s = new_slide('Dữ liệu: 387 cell của bốn bộ công khai và 18 cell NASA')
chk('22.431', '56.697', '82.450', '144.224', '305.802')
table(s, 0.58, 1.45, ['Bộ', 'Hoá học', 'Cell', 'Chu kỳ', 'Nhóm điều kiện', 'Tuổi thọ trung vị', 'Q danh định (Ah)'], [
    ['XJTU', 'NCM', '55', '22.431', '6', '392', '2,0'],
    ['TJU', 'NCA/NCM', '130', '56.697', '9', '377', '3,5 / 2,5'],
    ['MIT', 'LFP', '125', '82.450', '3', '622', '1,1'],
    ['HUST', 'LFP', '77', '144.224', '1', '1.855', '1,1'],
    ['NASA', 'NCA', '18', '1.597', '6', '–', '2,0'],
], [1.4, 1.65, 1.1, 1.6, 2.1, 2.15, 2.17], font=16, row_h=0.5)
body(s, 0.68, 4.75, 11.9, 1.6, [
    'Lọc ngoại lai nhân quả: quyết định giữ hay bỏ chu kỳ n chỉ dùng các chu kỳ trước n của chính cell; sau lọc còn '
    '305.802 chu kỳ ở bốn bộ chính',
    'Chuẩn hoá bằng thống kê của cell train; NASA được tự trích 16 đặc trưng sạc theo đúng định nghĩa của PINN4SOH'],
    size=16, bullets=True)
notes(s, 'Bốn bộ chính khác nhau rõ về hoá học và tuổi thọ. NASA nhỏ nhưng có hiện tượng hồi phục dung lượng sau khoảng '
         'nghỉ, nên được dùng để kiểm tra riêng ở cuối.')

# ═══════════════════════════════════════════════════════════ 5 · mô hình
s = new_slide('Mô hình đề xuất: PINN bán giám sát')
figure(s, 'fig_thesis_arch', 0.57, 1.3, 12.19, 4.05)
body(s, 0.68, 5.5, 11.9, 1.4, [
    'Mạng nghiệm Fφ dự đoán SOH; động học xám r = k(x)·g(u)·A(T) ≥ 0: tốc độ suy giảm không âm, tăng khi SOH giảm, '
    'theo Arrhenius',
    'Ba loss vật lý không dùng nhãn (phần dư ODE, đơn điệu, miền giá trị) áp lên mọi cell train, kể cả cell chưa đo'],
    size=15, bullets=True)
notes(s, 'Khi suy luận chỉ cần mạng nghiệm. Mạng tốc độ chỉ dùng lúc huấn luyện để tạo ràng buộc cho các cell chưa có nhãn.')

# ═══════════════════════════════════════════════════════════ 6 · hiện thực
s = new_slide('Hiện thực hệ thống')
figure(s, 'fig_thesis_pipeline', 0.57, 1.3, 12.19, 4.2)
nrun = chk('11.026')
body(s, 0.68, 5.65, 11.9, 1.3, [
    f'Python, PyTorch, scikit-learn; một lệnh chạy từ dữ liệu thô tới báo cáo; {nrun} lượt huấn luyện được trình bày',
    'Chạy lại 20 lượt mẫu trên máy khác cho dự đoán trùng từng bit; công cụ dòng lệnh ước lượng và dự báo cho một cell'],
    size=15, bullets=True)
notes(s, 'Năm khối: dữ liệu, mô hình, thí nghiệm, phân tích, báo cáo. Bộ lập lịch bỏ qua lượt đã có kết quả nên dừng giữa '
         'chừng rồi chạy lại vẫn tiếp tục được.')

# ═══════════════════════════════════════════════════════════ 7 · giao thức
s = new_slide('Giao thức đánh giá không rò rỉ')
figure(s, 'fig_thesis_leakage', 0.57, 1.3, 7.4, 5.4)
lo_hi = chk('1,2–1,8')
block(s, 8.25, 1.45, 4.5, [
    ('Chia theo cell', 'Kiểm định chéo 5 fold phân tầng theo chế độ thí nghiệm, 10 % cell validation, lặp 5 lần; mọi '
                       'mô hình dùng chung cách chia và cell có nhãn.'),
    ('Kiểm toán rò rỉ', f'Cùng một MLP, chia theo chu kỳ làm sai số thấp đi {lo_hi} lần; chép nhãn của chu kỳ train '
                        'gần nhất đã thấp hơn mọi mô hình được đánh giá đúng.'),
], size=15)
notes(s, 'Hình cho thấy cùng một mô hình nhưng đổi cách chia thì sai số đổi bao nhiêu. Vạch đứt là cách chia đúng, theo '
         'cell. Mọi kết quả sau đều dùng giao thức này.')

# ═══════════════════════════════════════════════════════════ 8 · theo lượng nhãn
s = new_slide('Kết quả 1: PINN bán giám sát theo lượng nhãn')
figure(s, 'fig_thesis_curve_physics', 0.57, 1.3, 7.6, 5.4)
chk('0,39', '0,90', '1,7--4,5', 'XJTU 0,96, TJU 1,09, MIT 0,85')
block(s, 8.45, 1.45, 4.3, [
    ('10 % nhãn', 'Sai số bằng 0,39–0,90 lần MLP, có ý nghĩa trên cả bốn bộ.'),
    ('Tiết kiệm nhãn', 'Với 10 % nhãn, PINN-semi đạt sai số mà MLP cần 17–45 % nhãn (1,7–4,5 lần).'),
    ('70 % nhãn', 'Tỉ số còn 0,85–1,09: khi đủ nhãn, vật lý ít thông tin mới.'),
], size=15)
notes(s, 'Đây là kết quả trung tâm: lợi ích của vật lý lớn nhất khi ít nhãn, vì loss vật lý buộc mạng cho quỹ đạo hợp lý '
         'trên chính các cell chưa có nhãn.')

# ═══════════════════════════════════════════════════════════ 9 · mô hình học sâu
s = new_slide('Kết quả 2: so với MLP, CNN, LSTM, Transformer')
figure(s, 'fig_models_compare', 0.57, 1.25, 12.19, 4.75)
chk('0,36--0,93', '42/48', '40/48', '0,77--0,96')
body(s, 0.68, 6.08, 11.4, 0.9, [
    'CNN, LSTM, Transformer có sai số bằng 0,36–0,93 lần MLP (có ý nghĩa ở 42/48 phép so) nhưng kém GBDT ở 40/48; '
    'chỉ tốt hơn GBDT trên HUST',
    'Mô hình cuối (PINN-lai, chọn lịch sử trên validation) có MAE bằng 0,77–0,96 lần mô hình chuỗi tốt nhất'],
    size=14, bullets=True)
notes(s, 'Vòng rỗng là 30 % nhãn, vòng đặc là 70 % nhãn. Mô hình chuỗi nhìn được nhiều chu kỳ nên tốt hơn MLP rõ rệt, '
         'nhưng GBDT vẫn mạnh hơn ở phần lớn trường hợp; mô hình cuối của đồ án thấp nhất ở mọi bộ.')

# ═══════════════════════════════════════════════════════════ 10 · bảng MAE, R²
s = new_slide('Bảng MAE và R² ở 30 % nhãn')
M = pd.read_csv(os.path.join(R, 'e30_metrics.csv'))
ORD = [('mlp', 'MLP'), ('cnn', 'CNN'), ('lstm', 'LSTM'), ('transformer', 'Transformer'), ('gbdt', 'GBDT'),
       ('pinn', 'PINN-semi'), ('pinn_lai', 'PINN-lai'), ('final', 'Mô hình cuối')]
DSL = ['XJTU', 'TJU', 'MIT', 'HUST']
g = M[M.f == 0.3].set_index(['dataset', 'model'])
num = lambda v, n: f'{v:.{n}f}'.replace('.', ',')
rows, best = [], []
for i, (m, lab) in enumerate(ORD):
    row = [lab]
    for ds in DSL:
        row += [num(1e3 * g.loc[(ds, m), 'MAE_cell'], 2), num(g.loc[(ds, m), 'R2'], 3)]
    rows.append(row)
for j, ds in enumerate(DSL):                     # tô màu giá trị tốt nhất mỗi cột
    mae = [g.loc[(ds, m), 'MAE_cell'] for m, _ in ORD]; r2 = [g.loc[(ds, m), 'R2'] for m, _ in ORD]
    for i in range(len(ORD)):
        if abs(mae[i] - min(mae)) < 1e-12:
            best.append((i, 1 + 2 * j))
        if abs(r2[i] - max(r2)) < 1e-12:
            best.append((i, 2 + 2 * j))
table(s, 0.58, 1.4, ['Mô hình'] + [x for ds in DSL for x in (f'{ds} MAE', 'R²')], rows,
      [2.05] + [1.265] * 8, font=14, row_h=0.47, best=best)
body(s, 0.68, 5.95, 11.9, 0.9, 'MAE theo cell (×10⁻³, thấp hơn là tốt hơn) và R² (gộp 5 fold của một lần lặp, bỏ 10 chu '
                               'kỳ đầu mỗi cell); 3 lần lặp × 5 fold, cùng fold và cell có nhãn. Màu xanh: tốt nhất mỗi cột.',
     size=13)
notes(s, 'Bảng R² mà đề cương cần, cùng MAE. R² gộp mọi cell nên phụ thuộc độ trải SOH của bộ; MAE theo cell là chỉ số chính.')

# ═══════════════════════════════════════════════════════════ 11 · quỹ đạo, vật lý
s = new_slide('Quỹ đạo dự đoán và tác động của vật lý')
figure(s, 'fig_pred_traj', 0.57, 1.25, 5.6, 5.55)
chk('13,6--31,1', '3,0--12,0', '0,1--7,7', '50--64', '62--95')
block(s, 6.55, 1.45, 6.2, [
    ('Tính đơn điệu', 'Ở 30 % nhãn, dự đoán của MLP tăng ngược ở 13,6–31,1 % số cặp chu kỳ, của PINN-semi chỉ '
                      '3,0–12,0 % (nhãn thật 0,1–7,7 %).'),
    ('Lợi ích đến từ đâu', 'Ở 10 % nhãn, 50–64 % mức giảm sai số đến từ loss vật lý trên cell không nhãn; ở 30 % nhãn, '
                           'riêng ràng buộc đơn điệu mang lại 62–95 %.'),
    ('Vật lý tác động qua đâu', 'Chủ yếu qua ràng buộc dọc quỹ đạo; tốc độ r học được chưa phải tốc độ suy giảm thật của '
                                'cell.'),
], size=15)
notes(s, 'Hình là cùng một cell test khi có 10, 30, 70 % nhãn. Đo trực tiếp trên mô hình đã huấn luyện cho thấy vật lý giúp '
         'chủ yếu qua tính đơn điệu dọc quỹ đạo.')

# ═══════════════════════════════════════════════════════════ 12 · PINN-lai
s = new_slide('Mô hình kết hợp PINN-lai')
chk('0,67--0,95', '10/16', '1,8--1,9')
body(s, 0.68, 1.3, 11.9, 0.9, 'ŷ = ω · (trung bình 3 PINN-semi) + (1 − ω) · GBDT, rồi lọc xu thế nhân quả trên w chu kỳ gần '
                              'nhất; ω và w chọn trên validation.', size=16)
figure(s, 'fig_mono_example', 0.57, 2.15, 12.19, 3.75)
body(s, 0.68, 6.02, 11.6, 0.95, [
    'Sai số bằng 0,67–0,95 lần PINN-semi (16/16 trường hợp); thấp hơn GBDT ở 10/16 trường hợp, 8 có ý nghĩa',
    'Bộ lọc xu thế làm trơn nhưng không bảo đảm SOH không tăng; ép không tăng làm sai số trên XJTU tăng 1,8–1,9 lần'],
    size=14, bullets=True)
notes(s, 'Phần lớn lợi ích đến từ việc ghép hai mô hình có lệch theo cell khác nhau. Hình bên trái là một cell XJTU có '
         'dung lượng tăng ở đầu đời: bộ lọc không tăng không theo được đoạn đó.')

# ═══════════════════════════════════════════════════════════ 13 · lịch sử
s = new_slide('Đầu vào có lịch sử của cell')
figure(s, 'fig_e29_lai', 0.57, 1.3, 12.19, 4.6)
chk('0,89--0,97', '16/16')
body(s, 0.68, 6.02, 11.6, 0.95, [
    'Trung bình trượt nhân quả của đặc trưng sạc: trên TJU và HUST từ 30 % nhãn, sai số còn 0,89–0,97 lần; trên XJTU và '
    'ở 10 % nhãn thì không giúp',
    'Chọn dùng hay không dùng lịch sử trên validation: PINN-lai thấp hơn GBDT ở 16/16 trường hợp, 13 có ý nghĩa'],
    size=14, bullets=True)
notes(s, 'Lựa chọn cuối dựa trên validation của đủ ba lần lặp, được đặt ra sau khi đã xem test, nên em trình bày là phân '
         'tích hậu kiểm.')

# ═══════════════════════════════════════════════════════════ 14 · PINN4SOH
s = new_slide('So với PINN4SOH trên đúng cách chia của bài')
figure(s, 'fig_e27_compare', 0.57, 1.25, 6.55, 5.55)
chk('0,23--0,63', '0,99', '1,65', '0,84--1,10', '0,33--0,94')
block(s, 7.45, 1.45, 5.3, [
    ('Tiền xử lý không rò rỉ', 'PINN-lai bằng 0,23–0,63 lần số công bố trên TJU và MIT, ngang trên XJTU (0,99), kém trên '
                               'HUST (1,65).'),
    ('Cùng tiền xử lý với bài', '0,84–1,10 lần ở chế độ đầy đủ, 0,33–0,94 lần ở chế độ ít mẫu, nhờ cell không nhãn đi '
                                'vào loss vật lý.'),
    ('Vì sao HUST kém', 'Chỉ số chu kỳ chuẩn hoá theo từng cell cho mô hình của bài biết phần đời đã qua.'),
], size=15)
notes(s, 'Cách chia được dựng lại từ nhật ký huấn luyện mà nhóm tác giả công bố. Cách tiền xử lý của bài có rò rỉ nên chỉ '
         'dùng để so cùng điều kiện với số công bố.')

# ═══════════════════════════════════════════════════════════ 15 · dự báo
s = new_slide('Dự báo SOH tương lai, dung lượng và RUL')
figure(s, 'fig_e22_traj', 0.57, 1.25, 7.25, 5.55)
chk('11,6--27,0', '0,21--0,60', '1,1--1,5', '78--84')
block(s, 8.1, 1.45, 4.65, [
    ('PINN-dự báo', 'Mạng trả tham số của nghiệm đóng của động học xám, kết hợp với GBDT, kèm khoảng dự báo 90 %.'),
    ('Khi cell không được đo', 'MAE của SOH trên toàn bộ phần đời còn lại 11,6–27,0 (×10⁻³), bằng 0,21–0,60 lần các mô '
                               'hình ngoại suy theo từng cell.'),
    ('RUL', 'GBDT không ngoại suy được ngoài tầm đã học nên sai số RUL gấp 1,1–1,5 lần; khoảng 90 % chỉ phủ 78–84 %.'),
], size=15)
notes(s, 'Đứng ở 40 % đời cell, mô hình dự báo phần còn lại; trục phải là dung lượng Ah. Khoảng dự báo còn hẹp hơn mức '
         'cần thiết, đây là hạn chế em ghi rõ trong báo cáo.')

# ═══════════════════════════════════════════════════════════ 16 · NASA dữ liệu
s = new_slide('Bộ NASA: dữ liệu và hồi phục sau khoảng nghỉ')
figure(s, 'fig_nasa_data', 0.57, 1.3, 12.19, 4.7)
chk('0,27', '0,40', '1,36')
body(s, 0.68, 6.1, 11.6, 0.85, [
    '18 cell NCA 2,0 Ah ở 4, 24 và 43 °C, sáu nhóm giao thức; SOH sau khoảng nghỉ dưới 10 giờ giảm trung vị 0,27 %, sau '
    '10–24 giờ tăng 0,40 %, từ 24 giờ trở lên tăng 1,36 %'], size=14, bullets=True)
notes(s, 'Hình phải là thay đổi SOH so với lần xả trước theo thời gian nghỉ: nghỉ càng lâu, dung lượng hồi phục càng nhiều. '
         'Đây là mẫu gây nhiễu mà mô hình phải xử lý.')

# ═══════════════════════════════════════════════════════════ 17 · NASA so sánh
s = new_slide('Bộ NASA: so với SambaMixer')
figure(s, 'fig_nasa_compare', 0.57, 1.3, 12.19, 4.7)
chk('6,17', '5,13', '4,12')
body(s, 0.68, 6.1, 11.6, 0.85, [
    'Kiểm định chéo: PINN-lai thấp nhất (6,17 % và 5,13 % SOH ở 30 và 70 % nhãn). Cách chia NASA-L: PINN-lai 4,12 %, '
    'SambaMixer công bố 0,96 % vì nó nhận đường xả của chính chu kỳ cần ước lượng'], size=14, bullets=True)
notes(s, 'Thanh gạch chéo là số công bố. SambaMixer dùng đường xả, nơi dung lượng gần như nằm sẵn trong đầu vào; đồ án '
         'chỉ dùng đoạn sạc trước lần xả, là dữ liệu BMS có ở mọi chu kỳ.')

# ═══════════════════════════════════════════════════════════ 18 · NASA quỹ đạo
s = new_slide('Bộ NASA: quỹ đạo trên ba cell đánh giá')
figure(s, 'fig_nasa_traj', 0.57, 1.3, 12.19, 4.7)
chk('101', '75--84', '1,02', '0,98')
body(s, 0.68, 6.1, 11.6, 0.85, [
    'B0006 sai nhiều nhất: bắt đầu ở 101 % SOH trong khi cell train cùng nhóm quanh 92 %, 75–84 % số chu kỳ nằm ngoài '
    'miền đặc trưng của cell train. Thêm giờ nghỉ làm đặc trưng không giúp (tỉ số 1,02 và 0,98)'], size=14, bullets=True)
notes(s, 'Ba cell đánh giá theo đúng cách chia của SambaMixer. Mô hình chuỗi tốt nhất trên cách chia này là CNN, vẽ cùng '
         'GBDT và PINN-lai.')

# ═══════════════════════════════════════════════════════════ 19 · NASA-L theo mức thông tin
s = new_slide('NASA-L: thêm thông tin để so với SambaMixer')
E31 = pd.read_csv(os.path.join(R, 'e31_tiers.csv'))
L31 = E31[(E31.split == 'samba') & (E31.ds == 'NASAfc')].set_index('tier')
TIERS = ['M0', 'q20', 'q10', 'q05', 'M2']
SAMBA = 0.961                                     # trung bình ba cell của SambaMixer-L (1,173; 1,197; 0,512)
rows, best = [], []
for i, (m, lab) in enumerate([('kalman', 'PINN-lai + Kalman'), ('kalman_k', 'PINN-lai + Kalman, κ (hậu kiểm)'),
                              ('capnhat', 'Mô hình cập nhật'), ('gan_nhat', 'Gần nhất'),
                              ('giu_nguyen', 'Giữ nguyên (không mô hình)')]):
    row = [lab]
    for j, t in enumerate(TIERS):
        v = 100 * float(L31.loc[t, m]) if (m == 'kalman' or t != 'M0') else None
        row.append('–' if v is None else chk(num(v, 2)))
        if v is not None and v < SAMBA:
            best.append((i, 1 + j))
    rows.append(row)
table(s, 0.58, 1.4, ['Cách ước lượng', 'Không đo', 'Mỗi 0,2L', 'Mỗi 0,1L', 'Mỗi 0,05L', 'Mọi lần xả trước'], rows,
      [2.7, 0.9, 0.9, 0.9, 0.95, 1.25], font=13, row_h=0.52, best=best)
body(s, 0.68, 4.75, 7.4, 1.4, 'MAE trung bình ba cell đánh giá (% SOH), 24 cột đầu vào, 3 hạt giống; ở mức đo thưa, sai '
                              'số tính trên các chu kỳ không được đo. Màu xanh: thấp hơn 0,96 % của SambaMixer.', size=13)
chk('4,12', '2,52', '2,16', '2,2', '0,96', '0,80', '0,63', '0,40')
block(s, 8.55, 1.4, 4.2, [
    ('Không đo cell', 'Thêm 7 đặc trưng của cả lần sạc: PINN-lai 4,12 → 2,52 %, LSTM 2,16 %, vẫn gấp 2,2 lần SambaMixer '
                      '(0,96 %, dùng chính lần xả).'),
    ('Đo thêm cell', 'Mô hình học cập nhật từ lần đo gần nhất: 0,63 % khi đo 8 lần mỗi cell, 0,40 % khi đo mỗi 0,1L, thấp '
                     'hơn SambaMixer trên cả ba cell. Bộ lọc Kalman chỉ thấp hơn khi biết mọi lần xả trước (0,80 %).'),
], size=14)
notes(s, 'SambaMixer nhận đường xả của chính chu kỳ cần ước lượng, nên nhãn gần như nằm trong đầu vào. Em tách ba mức thông '
         'tin về chính cell: không đo, đo thưa, và biết mọi lần xả trước. Không đo cell thì đồ án vẫn kém. Khi đo thêm '
         'vài lần, mô hình cập nhật học thẳng dung lượng đã đổi bao nhiêu kể từ lần đo gần nhất, từ thay đổi của đặc trưng '
         'sạc; với tám lần đo trong đời cell nó đã thấp hơn SambaMixer. Thiết kế của mô hình này chốt trước khi chạy, chọn '
         'chỉ trên các cell không thuộc tập đánh giá.')

# ═══════════════════════════════════════════════════════════ 20 · kiểm tra lại trên mọi bộ
s = new_slide('Kiểm tra lại trên mọi bộ: sai số theo lượng thông tin đo được')
figure(s, 'fig_e31_tiers', 0.57, 1.3, 12.19, 4.7)
chk('47/48', '48/48', '12/16')
body(s, 0.68, 6.1, 11.6, 0.85, [
    'Đo thưa: mô hình cập nhật thấp hơn bộ lọc Kalman ở 47/48 và thấp hơn giữ nguyên ở 48/48 trường hợp bộ × mức nhãn × '
    'lịch đo. Biết mọi lần xả trước: thấp hơn giữ nguyên ở 12/16'], size=14, bullets=True)
notes(s, 'Cùng các mức đo trên bốn bộ chính và trên kiểm định chéo của NASA. Đường tím là mô hình cập nhật, thấp nhất ở mọi '
         'ô. Bộ lọc Kalman chỉ sửa độ lệch của mô hình không đo, nên khi đã đo mọi chu kỳ trước nó còn kém cách giữ nguyên; '
         'mô hình cập nhật học thẳng thay đổi kể từ lần đo nên thắng cả hai.')

# ═══════════════════════════════════════════════════════════ 21 · ngân sách đo
s = new_slide('Ngân sách đo dung lượng')
figure(s, 'fig_e28_budget', 0.57, 1.25, 7.6, 5.55)
chk('1,2--6,0', '0,85--0,93', '0,97--1,05', '0,34--0,77')
block(s, 8.45, 1.45, 4.3, [
    ('Đo thưa mọi cell', 'Chỉ đo 1,2–6,0 % số chu kỳ, PINN-lai có sai số bằng 0,85–0,93 lần gắn nhãn đầy đủ cho 30 % số '
                         'cell và 0,97–1,05 lần khi gắn cho 70 %.'),
    ('Đo cả cell mới', 'Cập nhật bằng bộ lọc Kalman, sai số còn 0,34–0,77 lần GBDT học với 70 % cell có nhãn.'),
], size=15)
notes(s, 'Nên tính nhãn theo số lần đo chứ không theo số cell: rải vài lần đo lên mọi cell tốt hơn đo kín một phần cell.')

# ═══════════════════════════════════════════════════════════ 22 · đối chiếu đề cương
s = new_slide('Đối chiếu với đề cương (giai đoạn 1)')
table(s, 0.58, 1.4, ['Mục tiêu, nhiệm vụ trong đề cương', 'Đã thực hiện'], [
    ['Kiến thức nền, ràng buộc hoá lý', 'MLP, CNN, LSTM, Transformer, mã hoá vị trí, seq2seq; cơ chế suy giảm; PINN'],
    ['Thống kê mô tả, ngoại lai, mẫu nhiễu', 'Năm bộ dữ liệu; lọc ngoại lai nhân quả; hồi phục dung lượng của NASA'],
    ['So sánh ≥ 4 mô hình nền bằng MAE, RMSE, MAPE', 'Chín mô hình đối chứng; MAE, RMSE, MAPE, R²; kiểm định Wilcoxon'],
    ['Tinh chỉnh, giải thích', 'Độ dài cửa sổ; trọng số ghép, cửa sổ lọc; dạng lịch sử; đặc trưng cả lần sạc NASA'],
    ['Đề xuất cải tiến cho giai đoạn 2', 'PINN bán giám sát, PINN-lai, dự báo bằng nghiệm đóng của động học'],
    ['Phạm vi SOH, RUL; chỉ số; dự phòng', 'SOH trực tuyến; dự báo SOH, dung lượng, RUL; khoảng 90 %; cập nhật Kalman'],
    ['Bản mẫu, trực quan hoá', 'Đường ống tái lập được; công cụ dòng lệnh cho một cell; hình quỹ đạo'],
], [4.6, 7.57], font=14, row_h=0.6)
notes(s, 'Mọi mục của giai đoạn 1 đều đã làm. Điểm khác về hướng đi: GBDT mạnh hơn cả bốn mô hình học sâu, nên giai '
         'đoạn 2 cần giữ mô hình cây làm mốc so sánh.')

# ═══════════════════════════════════════════════════════════ 23 · kết luận
s = new_slide('Kết luận và hướng phát triển')
heading(s, 0.68, 1.45, 6.0, 'Kết luận', size=20)
body(s, 0.68, 1.95, 6.0, 4.6, [
    'PINN bán giám sát giảm sai số rõ nhất khi ít nhãn: 0,39–0,90 lần MLP ở 10 % nhãn',
    'Mô hình chuỗi tốt hơn MLP nhưng chưa vượt GBDT; GBDT là mốc so sánh cần có',
    'Kết hợp vật lý với GBDT (PINN-lai, mô hình cuối) cho sai số thấp nhất trong bảng tổng hợp',
    'Rải vài lần đo lên mọi cell hiệu quả hơn đo kín một phần cell',
    'NASA-L: đo thêm cell 8 lần, mô hình cập nhật cho 0,63 %, thấp hơn 0,96 % của SambaMixer'], size=16, bullets=True)
heading(s, 7.0, 1.45, 5.7, 'Giai đoạn 2', size=20)
body(s, 7.0, 1.95, 5.7, 4.6, [
    'Gắn ràng buộc vật lý vào mô hình chuỗi: LSTM, Transformer làm mạng nghiệm của PINN',
    'Khối pin nhiều cell: mô hình hoá mất cân bằng giữa các cell, chẳng hạn bằng GCN',
    'Giao diện theo dõi dựng từ công cụ dòng lệnh hiện có',
    'Lịch đo chủ động theo độ bất định của mô hình'], size=16, bullets=True)
fin_best = all(min(M[(M.dataset == ds) & (M.f == fr)].MAE_cell) >= float(
    M[(M.dataset == ds) & (M.f == fr) & (M.model == 'final')].MAE_cell.iloc[0]) - 1e-12
    for ds in DSL for fr in sorted(M.f.unique()))
assert fin_best                                   # "sai số thấp nhất trong bảng tổng hợp"
notes(s, 'Tóm lại: vật lý giúp nhiều nhất khi nhãn ít, mô hình cây vẫn là đối thủ khó nhất, và ghép hai thứ lại cho kết '
         'quả tốt nhất. Giai đoạn 2 chuyển sang khối pin nhiều cell.')

# ═══════════════════════════════════════════════════════════ 24 · cảm ơn
s = T.dup_slide(prs, COVER)
for sh in list(s.shapes):
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith('ĐỒ ÁN'):
        continue
    if t.startswith('Dự đoán'):
        T.set_text(sh, 'Cảm ơn Thầy Cô đã lắng nghe')
    elif t.startswith('Huỳnh Quốc Thắng'):
        T.set_text(sh, 'Em xin nhận câu hỏi')
    else:
        T.remove_shape(sh)
notes(s, 'Em xin cảm ơn và xin nhận câu hỏi.')

# ─────────────────────────────────────────────────────────────── dọn và đánh số
for i in range(n_old - 1, 0, -1):                  # bỏ các slide cũ, giữ bìa
    T.delete_slide(prs, i)
last = len(prs.slides) - 1
for i, s in enumerate(prs.slides):
    if 0 < i < last:
        T.set_text(T.page_shape(s), f'{i + 1:02d}')
prs.save(DST)
print('đã ghi', DST, f'({len(prs.slides)} slide)')
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', HERE, DST], check=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print('đã ghi', DST.replace('.pptx', '.pdf'))
