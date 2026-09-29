"""
Công cụ nhỏ để cập nhật bộ slide báo cáo (khuôn BK, 13,33 × 7,5 in) bằng python-pptx mà vẫn giữ
nguyên định dạng gốc: nhân bản slide, điền bảng, thay chữ giữ kiểu chữ, chèn hình vừa khung.

python-pptx không có hàm nhân bản slide; `dup_slide` chép cây shape và các quan hệ ảnh (logo),
bỏ qua biểu đồ (biểu đồ nhân bản sẽ dùng chung dữ liệu với slide gốc).
"""
from __future__ import annotations
import copy
from typing import List, Sequence
from lxml import etree
from pptx.util import Emu, Inches, Pt
from pptx.oxml.ns import qn

IMG_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'


# ─────────────────────────────────────────────────────────── slide
def dup_slide(prs, src_index: int):
    """Nhân bản slide src_index (0-based) vào cuối bộ. Trả về slide mới."""
    src = prs.slides[src_index]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):                      # bỏ placeholder mặc định của layout
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype == IMG_REL:
            rid_map[rid] = new.part.relate_to(rel._target, IMG_REL)
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        if el.tag == qn('p:graphicFrame') and el.find('.//' + qn('c:chart')) is not None:
            continue                                  # không chép biểu đồ
        e = copy.deepcopy(el)
        for node in e.iter():
            for attr in (qn('r:embed'), qn('r:link'), qn('r:id')):
                v = node.get(attr)
                if v in rid_map:
                    node.set(attr, rid_map[v])
        new.shapes._spTree.append(e)
    bg = src._element.find(qn('p:cSld')).find(qn('p:bg'))
    if bg is not None:
        new._element.find(qn('p:cSld')).insert(0, copy.deepcopy(bg))
    return new


def move_slide(prs, old: int, new: int):
    lst = prs.slides._sldIdLst
    el = list(lst)[old]
    lst.remove(el)
    lst.insert(new, el)


def delete_slide(prs, index: int):
    lst = prs.slides._sldIdLst
    el = list(lst)[index]
    prs.part.drop_rel(el.get(qn('r:id')))
    lst.remove(el)


# ─────────────────────────────────────────────────────────── shape
def shapes_by_text(slide, prefix: str):
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith(prefix)]


def title_shape(slide):
    """Tiêu đề = khối chữ nằm trong dải xanh đầu slide (top < 0,5 in)."""
    c = [s for s in slide.shapes if s.has_text_frame and s.top < Inches(0.5) and s.text_frame.text.strip()]
    return c[0]


def footnote_shape(slide):
    c = [s for s in slide.shapes if s.has_text_frame and Inches(6.5) < s.top < Inches(6.9)]
    return c[0] if c else None


def page_shape(slide):
    c = [s for s in slide.shapes if s.has_text_frame and s.top >= Inches(6.9) and s.left > Inches(11.5)]
    return c[0] if c else None


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_text(shape, text):
    """Thay chữ, giữ kiểu của run đầu tiên. `text` là chuỗi (xuống dòng = đoạn mới) hoặc list đoạn;
    mỗi đoạn là chuỗi hoặc list (chuỗi, đậm?) để trộn đậm/thường trong một đoạn."""
    tf = shape.text_frame
    paras = text.split('\n') if isinstance(text, str) else list(text)
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else None
    rpr = copy.deepcopy(r0._r.find(qn('a:rPr'))) if r0 is not None else None
    ppr = copy.deepcopy(p0._p.find(qn('a:pPr')))
    txBody = tf._txBody
    for p in list(txBody.findall(qn('a:p'))):
        txBody.remove(p)
    for para in paras:
        p = etree.SubElement(txBody, qn('a:p'))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        segs = [(para, None)] if isinstance(para, str) else para
        for seg, bold in segs:
            r = etree.SubElement(p, qn('a:r'))
            if rpr is not None:
                rp = copy.deepcopy(rpr)
                if bold is not None:
                    rp.set('b', '1' if bold else '0')
                r.append(rp)
            t = etree.SubElement(r, qn('a:t'))
            t.text = seg
    return shape


def set_font_size(shape, pt: float):
    for r in shape._element.iter(qn('a:rPr')):
        r.set('sz', str(int(round(pt * 100))))
    for r in shape._element.iter(qn('a:defRPr')):
        r.set('sz', str(int(round(pt * 100))))


def set_color(shape, hex6: str):
    for rpr in shape._element.iter(qn('a:rPr')):
        f = rpr.find(qn('a:solidFill'))
        if f is not None:
            c = f.find(qn('a:srgbClr'))
            if c is not None:
                c.set('val', hex6)


def add_picture_fit(slide, path: str, left, top, width, height, align='center'):
    """Chèn ảnh vừa khung (giữ tỉ lệ), căn giữa khung."""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    s = min(width / w, height / h)
    pw, ph = int(w * s), int(h * s)
    x = left + (width - pw) // 2 if align == 'center' else left
    y = top + (height - ph) // 2
    return slide.shapes.add_picture(path, x, y, pw, ph)


def clone_shape(src_shape, dst_slide):
    e = copy.deepcopy(src_shape._element)
    dst_slide.shapes._spTree.append(e)
    return dst_slide.shapes[-1]


# ─────────────────────────────────────────────────────────── table
def _set_cell(cell, text: str, bold=None, color=None):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    if not runs:
        p.add_run()
        runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    rpr = runs[0]._r.find(qn('a:rPr'))
    if rpr is not None and bold is not None:
        rpr.set('b', '1' if bold else '0')
    if color is not None and rpr is not None:
        f = rpr.find(qn('a:solidFill'))
        if f is not None and f.find(qn('a:srgbClr')) is not None:
            f.find(qn('a:srgbClr')).set('val', color)


def fill_table(shape, rows: Sequence[Sequence[str]], header: Sequence[str] = None,
               col_widths_in: Sequence[float] = None, font_pt: float = None, row_h_in: float = None,
               bold_cols=(), colors: dict = None):
    """Điền bảng: đổi số hàng (nhân bản hàng thân theo đúng nhịp màu xen kẽ), số cột giữ nguyên
    trừ khi col_widths_in cho số cột khác (khi đó nhân bản/xoá cột cuối). colors: {(r, c): hex}."""
    tbl = shape.table
    t = tbl._tbl
    # ---- cột
    if col_widths_in is not None:
        n_new = len(col_widths_in)
        grid = t.find(qn('a:tblGrid'))
        cols = grid.findall(qn('a:gridCol'))
        while len(cols) < n_new:
            grid.append(copy.deepcopy(cols[-1]))
            for tr in t.findall(qn('a:tr')):
                tr.append(copy.deepcopy(tr.findall(qn('a:tc'))[-1]))
            cols = grid.findall(qn('a:gridCol'))
        while len(cols) > n_new:
            grid.remove(cols[-1])
            for tr in t.findall(qn('a:tr')):
                tr.remove(tr.findall(qn('a:tc'))[-1])
            cols = grid.findall(qn('a:gridCol'))
        for c, w in zip(cols, col_widths_in):
            c.set('w', str(int(Inches(w))))
    # ---- hàng
    trs = t.findall(qn('a:tr'))
    body_tpl = [copy.deepcopy(trs[1]), copy.deepcopy(trs[2] if len(trs) > 2 else trs[1])]
    for tr in trs[1:]:
        t.remove(tr)
    for i in range(len(rows)):
        t.append(copy.deepcopy(body_tpl[i % 2]))
    trs = t.findall(qn('a:tr'))
    if row_h_in is not None:
        for tr in trs:
            tr.set('h', str(int(Inches(row_h_in))))
    if header is not None:
        for j, h in enumerate(header):
            _set_cell(tbl.cell(0, j), h)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            col = (colors or {}).get((i - 1, j))
            _set_cell(tbl.cell(i, j), str(v), bold=(True if j in bold_cols else None), color=col)
    if font_pt is not None:
        for rpr in t.iter(qn('a:rPr')):
            rpr.set('sz', str(int(round(font_pt * 100))))
    total_h = sum(int(tr.get('h')) for tr in t.findall(qn('a:tr')))
    shape.height = Emu(total_h)
    if col_widths_in is not None:
        shape.width = Emu(int(sum(Inches(w) for w in col_widths_in)))
    return shape


def set_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text
