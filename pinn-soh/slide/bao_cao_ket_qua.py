"""
Bộ slide báo cáo kết quả (bản mới) — dựng trên khuôn BK của bộ slide 29/09.

    python slide/bao_cao_ket_qua.py <bản-29-09.pptx> <bản-mới.pptx>

Chạy sau analyze_e13.py. Chỉ giữ phần mô hình và kết quả mới; chữ trên slide tối thiểu,
hình không mang chú thích. Số liệu đọc từ results/.
"""
from __future__ import annotations
import os, sys
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); sys.path.insert(0, os.path.dirname(HERE))
import pptx_tools as T
import section_e13 as S

R = 'results'
DS = S.DS
f, pct = S.f, S.pct
NAVY, INK, GREEN = '20567B', '163047', '0D6E56'


def build(src, dst):
    d = S.load(); F = S.facts(d); tests = d['tests']
    prs = Presentation(src)
    n0 = len(prs.slides)
    TABLE_SRC, FRAME_SRC = 17, 8
    frame_shapes = [s for s in prs.slides[FRAME_SRC].shapes
                    if not s.has_chart and s.left < Inches(1) and Inches(1.3) < s.top < Inches(1.7)]

    def drop_foot(s):
        ft = T.footnote_shape(s)
        if ft is not None:
            T.remove_shape(ft)
        if s.has_notes_slide:
            s.notes_slide.notes_text_frame.text = ''

    def table_slide(title, header, rows, widths, font=15, row_h=0.5, msg=None, colors=None):
        s = T.dup_slide(prs, TABLE_SRC)
        T.set_text(T.title_shape(s), title)
        tb = [x for x in s.shapes if x.has_table][0]
        tb.top = Inches(1.6)
        T.fill_table(tb, rows, header=header, col_widths_in=widths, font_pt=font, row_h_in=row_h, colors=colors)
        box = [x for x in s.shapes if x.has_text_frame and Inches(5.5) < x.top < Inches(6.5)][0]
        if msg:
            T.set_text(box, msg)
            box.top = tb.top + tb.height + Inches(0.35)
        else:
            T.remove_shape(box)
        drop_foot(s)
        return s

    def figure_slide(title, img, side=None):
        s = T.dup_slide(prs, TABLE_SRC)
        T.set_text(T.title_shape(s), title)
        for x in [x for x in s.shapes if x.has_table or (x.has_text_frame and Inches(5.5) < x.top < Inches(6.5))]:
            T.remove_shape(x)
        drop_foot(s)
        fr = T.clone_shape(frame_shapes[0], s)
        fr.left, fr.top, fr.width, fr.height = Inches(0.57), Inches(1.4), Inches(12.19), Inches(5.3)
        w = Inches(11.9 if side is None else 8.3)
        T.add_picture_fit(s, img, Inches(0.72), Inches(1.55), w, Inches(5.0))
        if side:
            tb = s.shapes.add_textbox(Inches(9.3), Inches(2.3), Inches(3.3), Inches(3.6))
            tf = tb.text_frame; tf.word_wrap = True
            for i, (txt, bold) in enumerate(side):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                r = p.add_run(); r.text = txt
                r.font.size = Pt(17 if bold else 15); r.font.bold = bold; r.font.name = 'Arial'
                r.font.color.rgb = RGBColor.from_string(NAVY if bold else INK)
                p.space_after = Pt(6 if bold else 14)
        return s

    def text_slide(src_idx, title, blocks):
        s = prs.slides[src_idx]
        T.set_text(T.title_shape(s), title)
        zone = lambda x: x.has_text_frame and Inches(1.5) < x.top < Inches(6.3)
        heads = [x for x in s.shapes if zone(x) and x.height < Inches(0.5)]
        bodies = [x for x in s.shapes if zone(x) and x.height > Inches(0.8)]
        for (h, b), hs, bs in zip(blocks, heads, bodies):
            T.set_text(hs, h); T.set_text(bs, b)
        drop_foot(s)
        return s

    # ─────────────── giữ lại: bìa, bài toán, dữ liệu, mô hình, huấn luyện; sửa chữ
    s1 = prs.slides[0]
    T.set_text(T.shapes_by_text(s1, 'Báo cáo tiến độ')[0], 'Báo cáo kết quả\n30/09/2026')
    if s1.has_notes_slide:
        s1.notes_slide.notes_text_frame.text = ''

    text_slide(2, 'Bài toán', [
        ('Đầu vào và đầu ra', '16 đặc trưng của đoạn sạc CC–CV và chỉ số chu kỳ → SOH của chu kỳ đó.'),
        ('Câu hỏi 1', 'Cùng 30 % cell có nhãn, PINN có sai số thấp hơn MLP không?'),
        ('Câu hỏi 2', 'PINN dùng 30 % nhãn có thay được MLP dùng 70 % nhãn không?'),
    ])

    s5 = prs.slides[4]
    from pinnsoh.data import load_dataset
    rows = []
    info = {'XJTU': ('NCM', '25 °C'), 'TJU': ('NCA / NCM', '25 / 35 / 45 °C'),
            'MIT': ('LFP', '30 °C'), 'HUST': ('LFP', '30 °C')}
    tot = 0
    for ds in DS:
        cs = load_dataset('PINN4SOH/data', ds, clean='causal')
        n = sum(len(c.soh) for c in cs); tot += n
        rows.append([ds, str(len(cs)), f'{n:,}'.replace(',', '.'), *info[ds]])
    T.fill_table([x for x in s5.shapes if x.has_table][0], rows,
                 header=['Bộ dữ liệu', 'Cell', 'Số chu kỳ', 'Hoá học', 'Nhiệt độ'])
    T.set_text([x for x in s5.shapes if x.has_text_frame and Inches(5) < x.top < Inches(6)][0],
               f'387 cell · {tot:,} chu kỳ · 16 đặc trưng mỗi chu kỳ'.replace(',', '.'))
    drop_foot(s5)

    s7 = prs.slides[6]
    T.set_text(T.shapes_by_text(s7, 'Mạng nghiệm')[0], 'Mạng nghiệm Fφ\n17 → 128 → 128 → 64 → 1')
    drop_foot(s7)
    s8 = prs.slides[7]
    drop_foot(s8)
    keep = [0, 2, 4, 6, 7]

    # ─────────────── slide mới
    new = []
    new.append(table_slide('Thiết lập thí nghiệm', ['Hạng mục', 'Cách làm'], [
        ['Chia dữ liệu', 'Kiểm định chéo 5 fold theo cell, lặp 5 lần'],
        ['Mỗi fold', '70 % train · 10 % validation · 20 % test'],
        ['Mô hình so sánh', 'MLP, PINN-sup, PINN-semi — cùng mạng, cùng chuẩn hoá'],
        ['Lượng nhãn', '10 / 30 / 50 / 70 % số cell'],
        ['Chỉ số', 'MAE theo từng cell'],
    ], [3.0, 9.17], font=16, row_h=0.55))

    rows, colors = [], {}
    for ds in DS:
        r = S._row(tests, 'C vs A', ds)
        rows.append([ds, f(1e3 * r.mae_a, 2), f(1e3 * r.mae_b, 2), pct(1 - r.ti_so), pct(r.cell_b_tot_hon)])
        if r.p_holm < 0.05 and r.ti_so < 1:
            colors[(len(rows) - 1, 3)] = GREEN
    new.append(table_slide('Kết quả 1: cùng 30 % nhãn',
                           ['Bộ', 'MLP (×10⁻³)', 'PINN-semi (×10⁻³)', 'Giảm sai số', 'Cell PINN tốt hơn'],
                           rows, [2.2, 2.4, 2.6, 2.4, 2.57], msg=f'PINN-semi giảm sai số trên {len(F["h1_win"])}/4 bộ dữ liệu',
                           colors=colors))
    new.append(figure_slide('Tỉ số sai số giữa các mô hình', f'{R}/fig_E13_forest.png'))
    new.append(figure_slide('Sai số theo lượng nhãn', f'{R}/fig_E13_curve.png'))

    rows2 = []
    for ds in DS:
        r = F['h2'][ds]
        rows2.append([ds, f(1e3 * r.mae_a, 2), f(1e3 * r.mae_b, 2), f(r.ti_so),
                      'đạt' if ds in F['h2_ok'] else 'chưa đạt'])
    new.append(table_slide('Kết quả 2: PINN 30 % so với MLP 70 %',
                           ['Bộ', 'MLP 70 % (×10⁻³)', 'PINN 30 % (×10⁻³)', 'Tỉ số', 'Không kém hơn 10 %'],
                           rows2, [2.2, 2.6, 2.6, 1.8, 2.97],
                           msg=f'Đạt trên {", ".join(F["h2_ok"])}; chưa đạt trên '
                               f'{", ".join(ds for ds in DS if ds not in F["h2_ok"])}'))
    new.append(figure_slide('Sai số từng cell', f'{R}/fig_E13_cells.png'))
    new.append(figure_slide('Quỹ đạo SOH trên cell test', f'{R}/fig_E13_traj.png'))

    sec = d['sec']; mrows = []
    for ds in DS:
        for arm, name in [('A', 'MLP'), ('C', 'PINN-semi')]:
            g = sec[(sec.bo == ds) & (sec.nhanh == arm)].iloc[0]
            mrows.append([ds, name, f(1e3 * g.MAE_late, 2), f(g.EOL_MAE, 1), f(g.MonoViol100, 3)])
    new.append(table_slide('Cuối đời pin và độ trơn của dự đoán',
                           ['Bộ', 'Mô hình', 'MAE khi SOH ≤ 0,9 (×10⁻³)', 'Sai số EOL (chu kỳ)', 'Tăng ngược / 100 chu kỳ'],
                           mrows, [1.6, 2.2, 3.0, 2.6, 2.77], font=14, row_h=0.42))

    w = d['e14']
    erows = []
    for _, r in w.sort_values(['cap', 'k']).iterrows():
        erows.append([r.cap, str(int(r.k)), f(r.mlp, 4), f(r.pinn_semi, 4)])
    new.append(table_slide('Chuyển sang bộ dữ liệu khác',
                           ['Nguồn → đích', 'Cell đích có nhãn', 'MLP', 'PINN-semi'], erows,
                           [3.4, 2.8, 3.0, 2.97], font=14, row_h=0.42))

    demo = [p for p in sorted(os.listdir('demo')) if p.endswith('_soh.png')]
    if demo:
        new.append(figure_slide('Demo: ước lượng SOH cho một cell', os.path.join('demo', demo[0]), side=[
            ('Đầu vào', True), ('Dữ liệu sạc từng chu kỳ, không cần đo dung lượng', False),
            ('Đầu ra', True), ('SOH theo từng chu kỳ và cảnh báo khi xuống 0,85', False)]))

    lost2 = [ds for ds in DS if ds not in F['h2_ok']]
    text_slide(19, 'Kết luận', [
        ('Cùng lượng nhãn', f'PINN-semi có sai số thấp hơn MLP trên {len(F["h1_win"])}/4 bộ, rõ nhất trên MIT.'),
        ('Tiết kiệm nhãn', f'PINN 30 % nhãn đạt mức của MLP 70 % trên {", ".join(F["h2_ok"])}; '
                           f'chưa đạt trên {", ".join(lost2)}.'),
        ('Hướng tiếp theo', 'Bổ sung mô hình đối chứng khác, ước lượng độ bất định và cải thiện dự đoán cuối đời pin.'),
    ])

    # ─────────────── sắp xếp: giữ 5 slide gốc + slide mới + kết luận; xoá phần còn lại
    order_old = keep + [19]
    ids = list(prs.slides._sldIdLst)
    new_ids = ids[n0:]
    wanted = [ids[i] for i in keep] + new_ids + [ids[19]]
    for i in sorted(set(range(n0)) - set(order_old), reverse=True):
        T.delete_slide(prs, i)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        lst.remove(el)
    for el in wanted:
        lst.append(el)
    for i, s in enumerate(prs.slides, start=1):
        ps = T.page_shape(s)
        if ps is not None:
            T.set_text(ps, f'{i:02d}')
    prs.save(dst)
    print(f'đã ghi {dst}: {len(prs.slides)} slide')


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
